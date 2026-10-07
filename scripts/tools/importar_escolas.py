"""
Importa o arquivo escolas.xlsx para o banco de dados.
Mantem os dados existentes (oferta, lat/lng) e atualiza os cadastrais.
"""
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(RAIZ))

from encaminhamento.database import init_db, get_session
from encaminhamento.database.crud import create_school, get_school_by_name, update_school


def limpar(valor):
    """Remove preenchimentos vazios do tipo '------------'."""
    if valor is None:
        return ""
    texto = str(valor).strip()
    if not texto or set(texto) <= {"-", " "}:
        return ""
    return texto


def extrair_telefone(texto):
    """Pega o telefone do texto de contato (ex.: 'Maria 99865-5705')."""
    if not texto:
        return ""
    limpo = re.sub(r"\s+", " ", str(texto))
    # Formato mais comum na fonte: 4 ou 5 digitos, hifen, 4 digitos
    match = re.search(r"(\d{4,5})\s*-\s*(\d{4})", limpo)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    # Fallback: sequencia longa de digitos
    match = re.search(r"\b(\d{8,13})\b", limpo)
    return match.group(1) if match else ""


def separar_contato(texto):
    """
    Separa o contato em (nome, telefone).
    A fonte traz o telefone grudado no nome, ex.: 'Silvana Caldas 99982-6320'.
    """
    if not texto:
        return "", ""

    limpo = re.sub(r"\s+", " ", str(texto)).strip()

    telefone = extrair_telefone(limpo)
    nome = limpo
    if telefone:
        # Remove o telefone e qualquer separador que tenha sobrado
        nome = re.sub(r"[\s\-\.]*" + re.escape(telefone), "", nome)
        nome = re.sub(r"\s+", " ", nome).strip(" -.,")

    # Se sobrou so o telefone no contato, deixa o nome vazio
    if not any(c.isalpha() for c in nome):
        nome = ""

    return nome, telefone


def main():
    init_db()

    caminho = RAIZ / "data" / "escolas.xlsx"
    if not caminho.exists():
        print(f"Arquivo nao encontrado: {caminho}")
        print("Copie escolas.xlsx para a pasta data/ antes de rodar.")
        return

    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.active

    # Descobre a linha do cabeÃ§alho
    linha_cabecalho = None
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        celulas = [str(v).strip().upper() for v in row if v]
        if "ESCOLA" in celulas:
            linha_cabecalho = idx
            break

    if linha_cabecalho is None:
        print("Cabecalho nao encontrado em escolas.xlsx")
        return

    print(f"Cabecalho na linha {linha_cabecalho}")
    print(f"Total de linhas: {ws.max_row}\n")

    criadas = atualizadas = 0

    with get_session() as session:
        for row in ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
            numero = row[0]
            nome = limpar(row[1] if len(row) > 1 else None)

            # SÃ³ processa linhas numeradas com nome preenchido
            if not nome or not isinstance(numero, (int, float)):
                continue

            modalidade = limpar(row[2] if len(row) > 2 else None)
            endereco = limpar(row[3] if len(row) > 3 else None)
            distrito = limpar(row[4] if len(row) > 4 else None)
            contato_bruto = limpar(row[5] if len(row) > 5 else None)
            telefone = limpar(row[7] if len(row) > 7 else None)
            email = limpar(row[8] if len(row) > 8 else None)

            # O contato vem com o telefone grudado: separa em nome e telefone
            nome_contato, telefone_do_contato = separar_contato(contato_bruto)

            if not telefone:
                telefone = telefone_do_contato

            endereco_completo = f"{endereco}, {distrito} Distrito" if endereco and distrito else endereco

            existente = get_school_by_name(session, nome)

            if existente:
                # Preserva oferta e coordenadas ja cadastradas
                update_school(
                    session, existente.id,
                    address=endereco_completo,
                    distrito=distrito,
                    modalidade=modalidade,
                    contato=nome_contato,
                    phone=telefone,
                    email=email,
                )
                atualizadas += 1
            else:
                create_school(
                    session, nome,
                    address=endereco_completo,
                    distrito=distrito,
                    modalidade=modalidade,
                    contato=nome_contato,
                    phone=telefone,
                    email=email,
                    is_origin=True,
                    is_destination=True,
                    oferta=0,
                )
                criadas += 1

    print(f"{criadas} escola(s) cadastrada(s)")
    print(f"{atualizadas} escola(s) atualizada(s)")

    # Resumo do que ainda falta preencher
    from encaminhamento.database.crud import list_schools
    with get_session() as session:
        escolas = list_schools(session)
        sem_oferta = [e for e in escolas if not e.oferta]
        sem_coord = [e for e in escolas if e.latitude is None]

    print(f"\nPendencias:")
    print(f"  Sem capacidade definida: {len(sem_oferta)}")
    print(f"  Sem geolocalizacao:      {len(sem_coord)}")


if __name__ == "__main__":
    main()
