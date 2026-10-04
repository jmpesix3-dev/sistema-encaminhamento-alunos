"""
Geracao de planilhas no formato do encaminhamento.

Formato modelodownload:
  Linha 1: Unidade Escolar de destino, Turma, Nº de alunos encaminhados
  Linha 2: QUADRO DE ENCAMINHAMENTO DE ALUNOS
  Linha 3: Nº | NOME DO ALUNO | ENDEREÇO | UNIDADE DE ORIGEM
"""
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from encaminhamento.config import EXCEL_EXPORT_DIR
from encaminhamento.database.models import School, Student
from encaminhamento.utils.arquivos import nome_seguro

ANO_PADRAO = 2025

CABECALHOS = ["Nº", "NOME DO ALUNO", "ENDEREÇO", "UNIDADE DE ORIGEM"]
LARGURAS = {"A": 8, "B": 40, "C": 55, "D": 45}

STYLE_TITULO = Font(bold=True, size=14)
STYLE_SUBTITULO = Font(bold=True, size=12)
STYLE_HEADER = Font(bold=True, size=11, color="FFFFFF")
STYLE_NORMAL = Font(size=11)
FILL_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
ALINHA_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALINHA_ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)
BORDA = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def export_students_to_excel(
    students: List[Student],
    destination_school: School,
    origin_school: School,
    class_name: str = "",
    output_path: str = None,
    ano: int = ANO_PADRAO,
) -> str:
    """Grava a lista de alunos no formato modelodownload."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "modelodownload"

    total = len(students)

    ws.merge_cells("A1:D1")
    ws["A1"] = (
        f"Unidade Escolar de destino: {destination_school.name}"
        f"          Turma: {class_name}"
        f"          Nº de alunos encaminhados: {total}      / Oferta 20"
    )
    ws["A1"].font = STYLE_TITULO
    ws["A1"].alignment = ALINHA_ESQUERDA

    ws.merge_cells("A2:D2")
    ws["A2"] = f"QUADRO DE ENCAMINHAMENTO DE ALUNOS {ano}"
    ws["A2"].font = STYLE_SUBTITULO
    ws["A2"].alignment = ALINHA_CENTRO

    for col, titulo in enumerate(CABECALHOS, 1):
        c = ws.cell(row=3, column=col, value=titulo)
        c.font = STYLE_HEADER
        c.fill = FILL_HEADER
        c.alignment = ALINHA_CENTRO
        c.border = BORDA

    for i, aluno in enumerate(students, 1):
        linha = 3 + i
        for col, valor, alinhamento in (
            (1, i, ALINHA_CENTRO),
            (2, aluno.name, ALINHA_ESQUERDA),
            (3, aluno.full_address, ALINHA_ESQUERDA),
            (4, origin_school.name, ALINHA_ESQUERDA),
        ):
            c = ws.cell(row=linha, column=col, value=valor)
            c.font = STYLE_NORMAL
            c.alignment = alinhamento
            c.border = BORDA

    for coluna, largura in LARGURAS.items():
        ws.column_dimensions[coluna].width = largura

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[3].height = 25
    for i in range(4, 4 + total):
        ws.row_dimensions[i].height = 22

    if output_path is None:
        output_path = (
            EXCEL_EXPORT_DIR / (
            f"encaminhamento_{nome_seguro(destination_school.name, 'escola')}"
            f"_{total}alunos.xlsx"
        )
        )

    wb.save(output_path)
    return str(output_path)


def create_template_modeloupload(output_path: str = None, ano: int = ANO_PADRAO) -> str:
    """Cria o modelo em branco para as escolas preencherem."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "modeloupload"

    ws.merge_cells("A1:E1")
    ws["A1"] = f"ENCAMINHAMENTO DE ALUNOS PARA {ano}"
    ws["A1"].font = STYLE_TITULO
    ws["A1"].alignment = ALINHA_CENTRO
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:E2")
    ws["A2"] = "Unidade Escolar de Origem: " + "_" * 48
    ws["A2"].font = Font(bold=True, size=11)
    ws.row_dimensions[2].height = 25

    ws.merge_cells("A3:E3")
    ws["A3"] = "Turma: __________     Total de alunos encaminhados: ______"
    ws["A3"].font = Font(bold=True, size=11)
    ws.row_dimensions[3].height = 25

    cabecalhos = [
        "Nº",
        "NOME DO ALUNO (completo)",
        "ENDEREÇO COMPLETO E COM PONTO DE REFERÊNCIA",
        "ESCOLA DE DESTINO",
        "ASSINATURA DO RESPONSÁVEL LEGAL",
    ]
    for col, titulo in enumerate(cabecalhos, 1):
        c = ws.cell(row=4, column=col, value=titulo)
        c.font = STYLE_HEADER
        c.fill = FILL_HEADER
        c.alignment = ALINHA_CENTRO
        c.border = BORDA
    ws.row_dimensions[4].height = 30

    # 30 alunos de exemplo, cada um com a linha da 2a opcao
    for i in range(1, 31):
        linha = 4 + (i * 2) - 1
        linha2 = linha + 1

        c = ws.cell(row=linha, column=1, value=i)
        c.font = STYLE_NORMAL
        c.alignment = ALINHA_CENTRO
        c.border = BORDA

        c = ws.cell(row=linha, column=2, value="Nome Civil: ")
        c.font = STYLE_NORMAL
        c.alignment = ALINHA_ESQUERDA
        c.border = BORDA

        ws.cell(row=linha, column=3).border = BORDA

        c = ws.cell(row=linha, column=4, value="1ª")
        c.font = STYLE_NORMAL
        c.alignment = ALINHA_CENTRO
        c.border = BORDA

        ws.cell(row=linha, column=5).border = BORDA

        c = ws.cell(row=linha2, column=4, value="2ª")
        c.font = STYLE_NORMAL
        c.alignment = ALINHA_CENTRO
        c.border = BORDA
        for col in (1, 2, 3, 5):
            ws.cell(row=linha2, column=col).border = BORDA
        ws.row_dimensions[linha2].height = 20

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 30

    rodape = 4 + 30 * 2 + 1
    ws.merge_cells(f"A{rodape}:E{rodape}")
    ws[f"A{rodape}"] = "Data: _____/_____/______."
    ws[f"A{rodape}"].font = Font(bold=True)

    rodape += 1
    ws.merge_cells(f"A{rodape}:E{rodape}")
    ws[f"A{rodape}"] = (
        "Secretário" + " " * 90 +
        "Diretor ou Responsável pela Unidade de Ensino"
    )

    if output_path is None:
        output_path = EXCEL_EXPORT_DIR / "modeloupload_template.xlsx"

    wb.save(output_path)
    return str(output_path)
